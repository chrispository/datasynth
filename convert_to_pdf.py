import os
import argparse
import io
import re
from email import policy
from email.parser import BytesParser
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from pypdf import PdfWriter, PdfReader

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

try:
    from docx import Document

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def parse_eml(file_path):
    with open(file_path, "rb") as f:
        msg = BytesParser(policy=policy.default).parse(f)

    subject = msg.get("subject", "(No Subject)")
    from_ = msg.get("from", "(Unknown Sender)")
    to = msg.get("to", "(Unknown Recipient)")
    date = msg.get("date", "(Unknown Date)")

    body = ""
    if msg.is_multipart():
        for part in msg.walk():
            ctype = part.get_content_type()
            cdispo = str(part.get("Content-Disposition"))
            if ctype == "text/plain" and "attachment" not in cdispo:
                body = part.get_content()
                break
    else:
        body = msg.get_content()

    return {"subject": subject, "from": from_, "to": to, "sent": date, "body": body}


def sanitize_text(text, collapse_whitespace=True):
    if not text:
        return ""
    if not isinstance(text, str):
        text = str(text)

    if collapse_whitespace:
        # Normalize whitespace (handles folded RFC 5322 headers)
        text = re.sub(r"\s+", " ", text).strip()
    
    # We are now using a Unicode font (DejaVu), so we don't need to aggressively 
    # downgrade to Latin-1. We just return the text.
    return text


from pdf_utils import init_pdf

class PDFConverter:
    def __init__(self, output_path):
        self.output_path = output_path

        # Use shared initialization
        self.pdf = init_pdf()

    def ensure_page(self):
        if self.pdf.page_no() == 0:
            self.pdf.add_page()
            self.pdf.set_font("DejaVu", size=10)

    def _render_email_headers(self, headers):
        """Render email headers with label bold, value regular."""
        for label, value in headers:
            self.pdf.set_font("DejaVu", "B", 10)
            label_width = self.pdf.get_string_width(label) + 2
            self.pdf.cell(
                label_width, 6, text=label, new_x=XPos.RIGHT, new_y=YPos.TOP
            )

            self.pdf.set_font("DejaVu", "", 10)
            self.pdf.multi_cell(
                0,
                6,
                text=sanitize_text(value),
                new_x=XPos.LMARGIN,
                new_y=YPos.NEXT,
            )

    def _render_separator_line(self):
        """Render a horizontal separator line."""
        self.pdf.ln(3)
        self.pdf.set_draw_color(180, 180, 180)
        y = self.pdf.get_y()
        self.pdf.line(10, y, 200, y)
        self.pdf.ln(5)
        self.pdf.set_draw_color(0, 0, 0)

    def add_eml(self, file_path):
        data = parse_eml(file_path)

        # Start each email on a new page
        self.pdf.add_page()
        self.pdf.set_font("DejaVu", size=10)

        # Headers - label bold, value regular (like real email clients)
        headers = [
            ("From:", data["from"]),
            ("Sent:", data["sent"]),
            ("To:", data["to"]),
            ("Subject:", data["subject"]),
        ]
        self._render_email_headers(headers)
        self._render_separator_line()

        # Body
        self.pdf.set_font("DejaVu", "", 10)
        body = data["body"].strip() if data["body"] else ""
        # Basic markdown-to-print cleanup
        body = body.replace("**", "")
        body = re.sub(r"^\* ", "• ", body, flags=re.MULTILINE)

        self.pdf.multi_cell(
            0,
            5,
            text=sanitize_text(body, collapse_whitespace=False),
            new_x=self.XPos.LMARGIN,
            new_y=self.YPos.NEXT,
        )

    def add_md(self, file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()

        # Parse headers
        headers = {}
        header_patterns = {
            "From:": r"^\*\*From:\*\* (.*)$",
            "Date:": r"^\*\*Date:\*\* (.*)$",
            "To:": r"^\*\*To:\*\* (.*)$",
            "Subject:": r"^\*\*Subject:\*\* (.*)$",
            "Cc:": r"^\*\*Cc:\*\* (.*)$",
            "Attachments:": r"^\*\*Attachments:\*\* (.*)$",
        }

        lines = content.split("\n")
        body_start_idx = 0
        for i, line in enumerate(lines):
            if line.strip() == "---":
                body_start_idx = i + 1
                break
            for label, pattern in header_patterns.items():
                match = re.match(pattern, line)
                if match:
                    headers[label] = match.group(1)

        body = "\n".join(lines[body_start_idx:]).strip()

        self.pdf.add_page()
        self.pdf.set_font("DejaVu", size=10)

        # Headers - label bold, value regular
        display_headers = [
            ("From:", headers.get("From:", "(Unknown Sender)")),
            (
                "Sent:",
                headers.get("Date:", "(Unknown Date)"),
            ),  # Use 'Sent:' for consistency
            ("To:", headers.get("To:", "(Unknown Recipient)")),
        ]
        if "Cc:" in headers:
            display_headers.append(("Cc:", headers["Cc:"]))

        display_headers.append(("Subject:", headers.get("Subject:", "(No Subject)")))
        
        # Attachments come after Subject
        if "Attachments:" in headers:
            display_headers.append(("Attachments:", headers["Attachments:"]))

        self._render_email_headers(display_headers)
        self._render_separator_line()

        self.pdf.set_font("DejaVu", "", 10)
        # Basic markdown-to-print cleanup
        body = body.replace("**", "")
        body = re.sub(r"^\* ", "• ", body, flags=re.MULTILINE)

        self.pdf.multi_cell(
            0,
            5,
            text=sanitize_text(body, collapse_whitespace=False),
            new_x=XPos.LMARGIN,
            new_y=YPos.NEXT,
        )

    def add_docx(self, file_path):
        self.pdf.add_page()

        if not HAS_DOCX:
            self.pdf.set_font("DejaVu", "I", 10)
            self.pdf.multi_cell(
                0,
                6,
                text=sanitize_text(
                    f"[Attachment: {os.path.basename(file_path)} - python-docx not installed]"
                ),
                new_x=self.XPos.LMARGIN,
                new_y=self.YPos.NEXT,
            )
            return

        try:
            doc = Document(file_path)

            # Distinctive header: Filename centered
            raw_title = os.path.basename(file_path)
            title = re.sub(r"^\d{4}_.*_attachment_", "", raw_title)
            title = title.replace(".docx", "").replace("_", " ")

            self.pdf.set_font("DejaVu", "B", 14)
            self.pdf.cell(
                0,
                10,
                text=sanitize_text(title),
                align="C",
                new_x=XPos.LMARGIN,
                new_y=YPos.NEXT,
            )
            self.pdf.ln(5)

            self.pdf.set_font("DejaVu", "", 11)
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    # Basic cleanup
                    clean_text = sanitize_text(text, collapse_whitespace=False)
                    clean_text = clean_text.replace("**", "")
                    self.pdf.multi_cell(
                        0,
                        5,
                        text=clean_text,
                        new_x=XPos.LMARGIN,
                        new_y=YPos.NEXT,
                    )
            self.pdf.ln(10)
        except Exception as e:
            self.pdf.set_font("DejaVu", "I", 10)
            self.pdf.multi_cell(
                0,
                6,
                text=sanitize_text(
                    f"[Error reading docx {os.path.basename(file_path)}: {e}]"
                ),
                new_x=self.XPos.LMARGIN,
                new_y=self.YPos.NEXT,
            )

    def save_temp_pdf(self):
        temp_path = "temp_generated.pdf"
        self.pdf.output(temp_path)
        return temp_path


def combine_files(folder_path, output_file):
    # Find all numbered files (matches 0001a_, 0001b_, etc.)
    files = []
    for f in os.listdir(folder_path):
        # Match pattern: NNNN followed by optional letter (a/b), then underscore
        match = re.match(r"^(\d{4}[a-z]?)_", f)
        if match:
            files.append((match.group(1), f))

    files.sort()  # Sort by prefix (0001, 0002...)

    writer = PdfWriter()
    converter = PDFConverter(output_file)

    for _, filename in files:
        file_path = os.path.join(folder_path, filename)
        ext = filename.lower().split(".")[-1]

        if ext == "eml":
            converter.add_eml(file_path)
        elif ext == "md":
            converter.add_md(file_path)
        elif ext == "docx":
            converter.add_docx(file_path)
        elif ext == "pdf":
            # Handle PDF merging
            # First, flush current converter to a temp PDF and add it to writer
            if len(converter.pdf.pages) > 0:
                temp_pdf = converter.save_temp_pdf()
                with open(temp_pdf, "rb") as f:
                    writer.append(f)
                os.remove(temp_pdf)
                # Reset converter for fresh start after this PDF
                converter = PDFConverter(output_file)

            # Now add the attachment PDF
            try:
                with open(file_path, "rb") as f:
                    writer.append(f)
            except Exception as e:
                print(f"Error merging PDF {filename}: {e}")

    # Final flush
    if len(converter.pdf.pages) > 0:
        temp_pdf = converter.save_temp_pdf()
        with open(temp_pdf, "rb") as f:
            writer.append(f)
        os.remove(temp_pdf)

    with open(output_file, "wb") as f:
        writer.write(f)
    print(f"Created combined PDF: {output_file}")


def convert_individual(folder_path):
    # This logic still exists if user doesn't want combine, but usually they do
    files = [f for f in os.listdir(folder_path) if f.lower().endswith((".eml", ".md"))]
    files.sort()
    for filename in files:
        file_path = os.path.join(folder_path, filename)
        output_name = os.path.splitext(filename)[0] + ".pdf"
        output_path = os.path.join(folder_path, output_name)

        converter = PDFConverter(output_path)
        ext = filename.lower().split(".")[-1]
        if ext == "eml":
            converter.add_eml(file_path)
        elif ext == "md":
            converter.add_md(file_path)
        converter.pdf.output(output_path)
        print(f"Converted: {output_name}")


def main():
    parser = argparse.ArgumentParser(description="Convert EML and attachments to PDF")
    parser.add_argument(
        "--folder", required=True, help="Path to folder containing files"
    )
    parser.add_argument(
        "--combine", action="store_true", help="Combine all into one PDF"
    )

    args = parser.parse_args()

    if not os.path.exists(args.folder):
        print(f"Error: Folder does not exist: {args.folder}")
        return

    if args.combine:
        folder_name = os.path.basename(os.path.normpath(args.folder))
        output_path = os.path.join(args.folder, f"{folder_name}_combined.pdf")
        combine_files(args.folder, output_path)
    else:
        convert_individual(args.folder)


if __name__ == "__main__":
    main()
